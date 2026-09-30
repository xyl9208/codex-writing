"""Export manuscript.md / manuscript_zh.md to a Word document (JTM submission style).

Figures (with their legends) and Table 1 are embedded in the main text right after the paragraph
that cites them for the first time; the separate 'Figure legends' / 'Tables' sections are therefore
not repeated at the end.  Usage: python 13_export_docx.py [input.md] [output.docx]
"""
import os, re, sys, pandas as pd
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MS = os.path.join(ROOT, "manuscript"); FIG = os.path.join(ROOT, "figures")
IN = sys.argv[1] if len(sys.argv) > 1 else "manuscript.md"; OUTNAME = sys.argv[2] if len(sys.argv) > 2 else "manuscript_JTM.docx"
ZH = IN.endswith("_zh.md")
lines = open(os.path.join(MS, IN), encoding="utf-8").read().split("\n")

FIGFILES = {1: ["Fig1_evidence_structure.png"], 2: ["Fig2_bulk_programmes.png"], 3: ["Fig3_sc_patient_level.png"], 4: ["Fig4_stage.png"], 5: ["Fig5_within_patient.png"], 6: ["Fig6_programme_organisation.png"]}
CITE_RE = re.compile(r"(?:Fig\.|Figs\.|图)\s?(\d)")
TABLE_RE = re.compile(r"(?:Table|表)\s?1\b"); TABLE2_RE = re.compile(r"(?:Table|表)\s?2\b")
SEC_LEGENDS = ("Figure legends", "图注"); SEC_TABLES = ("Tables", "表"); SEC_ADDITIONAL = ("Additional files", "附加文件")

# ---------------------------------------------------------------- collect legends; drop legend/table sections from the body
legends = {}; body = []; section = None
for line in lines:
    if line.startswith("## "): section = line[3:].strip()
    if section in SEC_LEGENDS:
        m = re.match(r"\*\*(?:Fig\.|图)\s?(\d)", line)
        if m: legends[int(m.group(1))] = line
        continue
    if section in SEC_TABLES: continue
    body.append(line)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
for s in doc.sections: s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.5)

def add_runs(par, text):
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|\^[^^]+\^)", text)
    for t in tokens:
        if not t: continue
        if t.startswith("**") and t.endswith("**"): r = par.add_run(t[2:-2]); r.bold = True
        elif t.startswith("*") and t.endswith("*"): r = par.add_run(t[1:-1]); r.italic = True
        elif t.startswith("^") and t.endswith("^"): r = par.add_run(t[1:-1]); r.font.superscript = True
        else: par.add_run(t)

def insert_figure(n):
    for f in FIGFILES.get(n, []):
        if not os.path.exists(os.path.join(FIG, f)): print("WARNING: missing figure file", f); continue
        doc.add_picture(os.path.join(FIG, f), width=Inches(6.3)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); add_runs(p, legends.get(n, f"**{'图' if ZH else 'Fig.'} {n}**")); p.paragraph_format.space_after = Pt(14)
    for r in p.runs: r.font.size = Pt(10)

def insert_table1():
    cap = doc.add_paragraph(); r = cap.add_run("表 1 " if ZH else "Table 1 "); r.bold = True
    cap.add_run("纳入研究的队列" if ZH else "Cohorts included in the study")
    t1 = pd.read_csv(os.path.join(MS, "Table1_cohorts.csv"))
    cols = ["Accession", "Disease", "Design", "Platform", "Cases", "Controls", "DE_method", "DEGs_FDR05", "Role"]
    table = doc.add_table(rows=1, cols=len(cols)); table.style = "Table Grid"
    for i, c in enumerate(cols): table.rows[0].cells[i].text = c.replace("_", " ")
    for _, r in t1.iterrows():
        cells = table.add_row().cells
        for i, c in enumerate(cols): cells[i].text = str(r[c])
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs: run.font.size = Pt(7.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def insert_table2():
    cap = doc.add_paragraph(); r = cap.add_run("表 2 " if ZH else "Table 2 "); r.bold = True
    cap.add_run("预设比较的证据矩阵（C、I、R；完整表见 Table2_evidence_matrix.csv）" if ZH else "Evidence matrix for the pre-specified comparisons (C, I and R; full table in Table2_evidence_matrix.csv)")
    t2 = pd.read_csv(os.path.join(MS, "Table2_evidence_matrix_main.csv")).fillna("")
    cols = ["block", "dataset", "comparison", "measure", "effect", "n", "inference", "estimation", "estimability"]
    hdr = ["数据块", "数据集", "比较", "指标", "效应 [95% CI]", "n", "推断", "估计", "可估计性"] if ZH else ["Block", "Dataset", "Comparison", "Measure", "Effect [95 % CI]", "n", "Inference", "Estimation", "Estimability"]
    table = doc.add_table(rows=1, cols=len(cols)); table.style = "Table Grid"
    for i, c in enumerate(hdr): table.rows[0].cells[i].text = c
    for _, r in t2.iterrows():
        cells = table.add_row().cells
        for i, c in enumerate(cols): cells[i].text = str(r[c])
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs: run.font.size = Pt(6.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

inserted = set(); table_done = False; table2_done = False; first = True; section = None
for line in body:
    line = line.rstrip()
    if not line or line == "---": continue
    if line.startswith("# ") and first:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(line[2:]); r.bold = True; r.font.size = Pt(15); first = False; continue
    if line.startswith("## "):
        section = line[3:].strip()
        if section in SEC_ADDITIONAL:  # figures never cited in the text are appended before the additional-files list
            for n in sorted(FIGFILES):
                if n not in inserted: insert_figure(n); inserted.add(n)
        doc.add_heading(section, level=1); continue
    if line.startswith("### "): doc.add_heading(line[4:], level=2); continue
    if re.match(r"^\d+\. ", line):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8); add_runs(p, line); continue
    p = doc.add_paragraph(); add_runs(p, line)
    if line.startswith("**Keywords") or line.startswith("**关键词"): p.paragraph_format.space_after = Pt(12)
    in_main = section not in (None, "Abstract", "摘要")
    if in_main and not table_done and TABLE_RE.search(line):
        insert_table1(); table_done = True
    if in_main and not table2_done and TABLE2_RE.search(line) and os.path.exists(os.path.join(MS, "Table2_evidence_matrix_main.csv")):
        insert_table2(); table2_done = True
    if in_main:
        for m in CITE_RE.finditer(line):
            n = int(m.group(1))
            if n in FIGFILES and n not in inserted: insert_figure(n); inserted.add(n)
out = os.path.join(MS, OUTNAME); doc.save(out); print("saved", out, "| figures embedded:", sorted(inserted), "| table1:", table_done, "| table2:", table2_done)
